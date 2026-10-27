from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from zipfile import ZipFile
import hashlib

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'output' / 'documents'
OUT.mkdir(parents=True, exist_ok=True)
PHOTO = Path(r'C:\Users\Pongo\OneDrive\Pictures\CV\foto cv.png')
d = Document()
s = d.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin = s.bottom_margin = Cm(1.6)
s.left_margin = s.right_margin = Cm(1.8)
normal = d.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0,0,0)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.space_after = Pt(4)
for border in list(d.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
for name in ['Title','Subtitle','Heading 1','Heading 2']:
    st=d.styles[name]; st.font.name='Arial'; st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(4)
d.styles['Title'].font.size=Pt(25)
d.styles['Title'].font.bold=True
d.styles['Heading 1'].font.size=Pt(11.5)
d.styles['Heading 1'].font.bold=True
d.styles['Heading 1'].paragraph_format.space_before=Pt(11)
d.styles['Heading 2'].font.size=Pt(10.5)
d.styles['Heading 2'].font.bold=True
d.styles['Heading 2'].paragraph_format.space_before=Pt(6)

def p(text='',bold=False):
    a=d.add_paragraph();a.add_run(text).bold=bold;return a
def link(a,label,url):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),a.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    c=OxmlElement('w:color');c.set(qn('w:val'),'17365D');pr.append(c)
    r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);a._p.append(h)
def heading(t): d.add_paragraph(t,'Heading 1')
def role(title,where,date):
    d.add_paragraph(title,'Heading 2')
    a=p(f'{where} | {date}');a.paragraph_format.space_after=Pt(3)
def bullet(t):
    a=d.add_paragraph(t,'List Bullet');a.paragraph_format.space_after=Pt(3)
    a.paragraph_format.line_spacing=1.05

# Embed the source photo byte for byte; float it beside the name and contact block.
a=d.add_paragraph('RAHMAWAN','Title')
run=a.add_run(); pic=run.add_picture(str(PHOTO),width=Cm(2.45))
inline=pic._inline
anchor=parse_xml(f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="114300" distR="0" simplePos="0" relativeHeight="251658240" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="0"><wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="margin"><wp:align>right</wp:align></wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV></wp:anchor>')
for tag in ['extent','effectExtent']:
    child=inline.find(qn('wp:'+tag))
    if child is not None: anchor.append(child)
anchor.append(OxmlElement('wp:wrapSquare'))
anchor[-1].set('wrapText','bothSides')
for tag in ['docPr','cNvGraphicFramePr']:
    child=inline.find(qn('wp:'+tag))
    if child is not None: anchor.append(child)
anchor.append(inline.find(qn('a:graphic')))
anchor.find(qn('wp:docPr')).set('descr','Original photograph of Rahmawan')
inline.getparent().replace(inline,anchor)
p('Management Graduate | Administration | Front End Development',True)
a=p();link(a,'+62 878-4311-0141','tel:+6287843110141');a.add_run(' | ');link(a,'rahm2668@gmail.com','mailto:rahm2668@gmail.com')
a=p();link(a,'linkedin.com/in/rahmawan-9a3395215','https://www.linkedin.com/in/rahmawan-9a3395215/');a.add_run(' | ');link(a,'github.com/rahm2668-cpu','https://github.com/rahm2668-cpu')
a=p();link(a,'Portfolio: portofolio-website-sigma-liart.vercel.app','https://portofolio-website-sigma-liart.vercel.app/')
heading('PROFILE')
p('Management graduate with experience in administration, document management, data processing, and operational support. Skilled in Microsoft Excel and Word, with a detail-oriented and organized approach to work. Gained front end development experience through internships and independent hospitality web projects using React and Next.js. Brings communication, coordination, and problem-solving skills to administrative, hospitality, and digital support roles.')
heading('WORK EXPERIENCE')
role('Front End Developer Intern','Codealpha','2026')
bullet('Completed assigned tasks according to project instructions and deadlines.')
bullet('Documented work and reported task progress while developing time management and teamwork skills.')
role('Front End Developer Project Based Virtual Intern','Core Initiative x Rakamin Academy | Depok','2025')
bullet('Participated in a structured, project-based front end development internship.')
bullet('Collaborated on assigned tasks, maintained work documentation, and reported progress against deadlines.')
role('Administration Staff','CV. Martiel Cahaya Perkasa | Bandung','2021 - 2022')
bullet('Managed administrative documents, prepared reports, and maintained an organized filing system.')
bullet('Entered, updated, and processed data using Microsoft Excel; checked completeness and accuracy before processing.')
bullet('Coordinated with internal teams to support day-to-day administrative operations.')
role('Part Time Driver','Grab','2019 - 2020')
bullet('Delivered customer service professionally and managed working time independently to meet daily targets.')
heading('EDUCATION')
role('Bachelor of Economics in Management','Universitas Garut','2018 - 2022')
p('GPA: 3.61')
p('Academic Achievement Scholarship (PPA) recipient, Ministry of Education, Culture, Research, and Technology.')

d.add_page_break()
heading('SELECTED PROJECTS')
d.add_paragraph('Kayarra Resort Hotel Website','Heading 2')
a=p();link(a,'kayarra-official-website.vercel.app','https://kayarra-official-website.vercel.app/')
bullet('Developed a responsive resort website presenting rooms, villas, facilities, photo galleries, and reservation inquiries.')
p('Technologies: Next.js, React.',True)
d.add_paragraph('OTA Management System Hotel','Heading 2')
a=p();link(a,'kayarra-ota-management-system.vercel.app','https://kayarra-ota-management-system.vercel.app/')
bullet('Built a hotel operations portfolio demo covering reservations, channel distribution, rate and inventory calendars, and revenue overviews using fictional sample data.')
p('Technologies: Next.js, React.',True)
heading('SKILLS AND TOOLS')
p('Administration: Microsoft Excel, Microsoft Word, data entry and processing, report preparation, document management, and operational support.')
p('Web development: HTML, CSS, React, Next.js, Tailwind CSS. Tools: Visual Studio Code.')
p('Professional skills: communication, coordination, teamwork, time management, attention to detail, and problem solving.')
p('English: B2 Upper Intermediate (EF SET, 2025).')
heading('CERTIFICATIONS')
for text in [
    'Core Initiative Frontend Project Based Internship Program | Core Initiative | 2026',
    'React Web Developer Expert | Dicoding | 2025',
    'Financial Literacy | Dicoding | 2025',
    'EF SET English Certificate, B2 Upper Intermediate | 2025',
    'Marketing | BNSP | 2021',
    'PTESOL | The Language Center, Indonesia University of Education | 2021',
]: bullet(text)
heading('ORGANIZATIONAL EXPERIENCE')
role('Ministry of External Relations','Badan Eksekutif Mahasiswa, Universitas Garut','2019 - 2020')
bullet('Supported internal and external organizational relations, coordination between student organizations, and administrative activities.')
role('Public Relations','BEM Sport and Art Competition Vol. 2','2019')
bullet('Communicated event information, responded to participant inquiries, and coordinated communication between the committee and external parties.')

d.core_properties.title='Rahmawan Curriculum Vitae'
d.core_properties.author='Rahmawan'
d.core_properties.subject='Professional experience, education, skills, and projects'
file=OUT/'Rahmawan-CV-Updated.docx'
d.save(file)
with ZipFile(file) as z:
    embedded=[z.read(n) for n in z.namelist() if n.startswith('word/media/')]
    assert any(hashlib.sha256(b).digest()==hashlib.sha256(PHOTO.read_bytes()).digest() for b in embedded)
    assert '<w:tbl>' not in z.read('word/document.xml').decode()
print(file)
print('Original photo bytes verified. Single column body, no tables.')
